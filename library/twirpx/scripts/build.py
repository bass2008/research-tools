import json, re, csv, collections, os
import openpyxl
from openpyxl.utils import get_column_letter
P = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "raw")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "twirpx-ezoterika-spisok")
items = {}
for l in open(f"{P}/items.jsonl", encoding="utf-8"):
    x = json.loads(l)
    path = x["sections"].replace("Раздел:", "").strip()
    if x["id"] in items:
        if path not in items[x["id"]]["paths"]: items[x["id"]]["paths"].append(path)
        continue
    x["paths"] = [path]; items[x["id"]] = x
rows = []
ISSUE = re.compile(r"\b(?:19|20)\d\d\s*№\s*\d")
dropped = [x for x in items.values() if x["page"].startswith("/files/nonscience/isoteric/periodic/") or ISSUE.search(x["title"])]
items = {k: x for k, x in items.items() if x not in dropped}
print("убрано журналов:", len(dropped), "из них вне раздела журналов:", sum(1 for x in dropped if not x["page"].startswith("/files/nonscience/isoteric/periodic/")))
status = {}
for l in open(f"{P}/status.jsonl", encoding="utf-8"):
    y = json.loads(l); status[y["id"]] = y["status"]
removed = [x for x in items.values() if status.get(x["id"]) == "copyright"]
unknown = [x for x in items.values() if x["id"] not in status]
items = {k: x for k, x in items.items() if status.get(k) == "approved"}
print("книг удалено правообладателем:", len(removed), "| без статуса (пропали с сайта):", len(unknown))
for x in items.values():
    d = x["description"]
    head = d[:160]
    year = re.search(r"\b(1[5-9]\d\d|20[0-2]\d)\b", head)
    pages = re.search(r"—\s*(\d+)\s*(?:с|p|pp|стр|S)\b", head)
    imprint = re.split(r"\s—\s", d, 1)[0] if " — " in head else ""
    path = x["paths"][0].replace("Вненаучные способы познания → ", "")
    rows.append([x["id"], x["title"], year.group(1) if year else "", imprint[:150], pages.group(1) if pages else "",
                 ", ".join(x["formats"]), x["size"], x["added"], path, "; ".join(p.replace("Вненаучные способы познания → ", "") for p in x["paths"][1:]),
                 f"https://www.twirpx.com/file/{x['id']}/"])
rows.sort(key=lambda r: (r[8], r[1].lower()))
head = ["ID", "Название (как на сайте)", "Год", "Выходные данные", "Страниц", "Форматы", "Размер", "Добавлен", "Раздел", "Ещё в разделах", "Ссылка"]
with open(OUT + ".csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f); w.writerow(head); w.writerows(rows)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Книги"
ws.append(head)
for r in rows: ws.append(r)
for i, wdt in enumerate([10, 70, 6, 45, 8, 12, 10, 11, 60, 30, 36], 1): ws.column_dimensions[get_column_letter(i)].width = wdt
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
ws3 = wb.create_sheet("Удалены правообладателем")
ws3.append(["ID", "Название (как на сайте)", "Раздел"])
for x in sorted(removed, key=lambda x: (x["paths"][0], x["title"].lower())):
    ws3.append([x["id"], x["title"], x["paths"][0].replace("Вненаучные способы познания → ", "")])
ws3.column_dimensions["B"].width = 70; ws3.column_dimensions["C"].width = 60
pages = list({json.loads(l)["page"]: json.loads(l) for l in open(f"{P}/pages.jsonl", encoding="utf-8")}.values())
by_sec = collections.Counter(r[8] for r in rows)
ws2 = wb.create_sheet("Разделы"); ws2.append(["Раздел", "Файлов прямо в разделе (собрано)", "По карте сайта (с подразделами)", "Адрес"])
for p in [p for p in pages if not p["page"].startswith("/files/nonscience/isoteric/periodic/")]: ws2.append([p["name"], p["items"], int(p["map_count"]) if p["map_count"] else None, "https://www.twirpx.com" + p["page"]])
ws2.column_dimensions["A"].width = 60; ws2.column_dimensions["D"].width = 70
wb.save(OUT + ".xlsx")
print("уникальных файлов:", len(rows), "| записей всего:", sum(1 for _ in open(f"{P}/items.jsonl", encoding="utf-8")), "| страниц:", len(pages))
print("с годом:", sum(1 for r in rows if r[2]), "| со страницами:", sum(1 for r in rows if r[4]))
print("в нескольких разделах:", sum(1 for r in rows if r[9]))
print("осталось с «Журналы» в разделе:", sum(1 for r in rows if "Журналы" in r[8]))
