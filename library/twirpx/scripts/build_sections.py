import json, re, csv, sys, os, collections
import openpyxl
from openpyxl.utils import get_column_letter
RAW, OUT, SECTIONS = sys.argv[1], sys.argv[2], sys.argv[3]
SKIP = [l.strip()[1:] for l in open(SECTIONS, encoding="utf-8") if l.startswith("-")]
ROOTS = [l.strip() for l in open(SECTIONS, encoding="utf-8") if l.strip() and not l.startswith("-")]
skipped_page = lambda u: any(u.startswith(p) for p in SKIP)
NOT_BOOK_TYPES = {"lp", "standard", "article", "paper", "guideline", "lecture", "package", "cheat", "control", "course",
                  "degree", "lab", "presentation", "ref", "report", "question", "software", "workbook"}
KIND = r"(?:Курсов\w*|Дипломн\w*|Реферат\w*|Контрольн\w*|Лекци\w*|Конспект\w*|Курс лекций|Шпаргал\w*|Презентаци\w*|Отч[её]т\w*|Ответы|Билеты|Тест\w*|Тренинг\w*|Доклад\w*|Эссе|Лабораторн\w*|Методическ\w*)"
NOT_BOOK_TITLE = re.compile(rf"^{KIND}\b|\b(?:Шпаргалк\w*|Семинар\w*|Учебно-методическ\w* комплекс\w*|Материалы для презентации|Лекци\w*|Лекция|рефераты|экзаменационн\w*|курсов\w* (?:работ\w*|проектирование)|дипломн\w* работ\w*|аналитическ\w* отч[её]т\w*|отч[её]т\w* по\b|Методические (?:указания|рекомендации)|Конспект лекций|Курс лекций|Курсовая работа|Дипломная работа|Контрольная работа|Реферат|Презентация)\b|^(?:Шаблон|Пример|Кейс\w*|Схема|Инструкция|Проект|Бриф\S*|Каталог\w*|Аналитический обзор|Исследовательская работа|Расч[её]тно.графическая работа|Маркетинговое исследование|Group research|Руководство - Brand\s?book)\b", re.I)
NOT_BOOK_DESC = re.compile(r"^(?:Каталог выставки|Конспект лекций|Курс лекций|Методическ\w+ (?:указани|рекомендаци)|Курсовая работа|Дипломная работа|Реферат|Презентация)", re.I)
ISSUE = re.compile(r"\b(?:19|20)\d\d\s*№\s*\d")
CONV_STRONG = re.compile(r"(?i)конверси|conversion|лендинг|landing page|a/b|split.?test|сплит-тест|юзабилити|usabilit|воронк\w* продаж|sales funnel|копирайт|copywrit|нейромаркет|neuromarket|checkout|\bcro\b")
CONV_TITLE = re.compile(r"(?i)убежд|persuas|influenc|психологи\w* (?:продаж|покупател|потребител|рекламы|влияния|цен)|pricing|продающ|оффер|e-?commerce|интернет-магазин|онлайн-продаж|customer experience|клиентск\w+ опыт|поведени\w+ (?:потребител|покупател)|consumer behavio|buyer|\bux\b|user experience")
SEO_SECTIONS = ("SEO и клиентская оптимизация", "Поисковые движки и механизмы")
SEO_TITLE = re.compile(r"(?i)\bseo\b|поисков\w* (?:оптимизац|продвижени|маркетинг)|продвижени\w* (?:сайт|в поисков)|search engine|\bserp\b|семантическ\w* ядр|ссылочн\w* (?:продвижени|масс)|link.?building|keyword research")
pages = [p for p in {json.loads(l)["page"]: json.loads(l) for l in open(f"{RAW}/pages.jsonl", encoding="utf-8")}.values() if not skipped_page(p["page"])]
items = {}
for l in open(f"{RAW}/items.jsonl", encoding="utf-8"):
    x = json.loads(l)
    if skipped_page(x["page"]): continue
    path = x["sections"].replace("Раздел:", "").strip()
    if x["id"] in items:
        if path not in items[x["id"]]["paths"]: items[x["id"]]["paths"].append(path)
        continue
    x["paths"] = [path]; items[x["id"]] = x
total = len(items)
journals = {k for k, x in items.items() if "/periodic/" in x["page"] or ISSUE.search(x["title"])}
removed = [x for k, x in items.items() if k not in journals and x["status"] == "copyright"]
other_status = collections.Counter(x["status"] for k, x in items.items() if k not in journals and x["status"] not in ("approved", "copyright"))
ok = [x for k, x in items.items() if k not in journals and x["status"] == "approved"]
def not_book(x):
    if x["file_type"] in NOT_BOOK_TYPES: return f"тип файла: {x['file_type']}"
    if NOT_BOOK_TITLE.search(x["title"]): return "название"
    if NOT_BOOK_DESC.search(x["description"]): return "описание"
    return ""
skipped = [(x, not_book(x)) for x in ok if not_book(x)]
books = [x for x in ok if not not_book(x)]
group_name = {p["page"]: p["name"] for p in pages if p["page"] in ROOTS}
group = lambda u: group_name[max((r for r in ROOTS if u.startswith(r)), key=len)]
short = lambda p: p.replace("Академическая и специальная литература → ", "").replace("Прикладная литература → ", "")
rows = []
for x in books:
    d = x["description"]; head = d[:160]
    year = re.search(r"\b(1[5-9]\d\d|20[0-2]\d)\b", head)
    pg = re.search(r"—\s*(\d+)\s*(?:с|p|pp|стр|S|р)\b", head)
    imprint = re.split(r"\s—\s", d, 1)[0] if " — " in head else ""
    lang = "рус" if re.search(r"[А-Яа-яЁё]", x["title"]) else "иностр"
    textbook = "да" if re.search(r"(?i)учебник|учебн\w* пособи|учебно-методическ|для студентов|для вузов", d[:250]) else ""
    conv = "да" if CONV_STRONG.search(x["title"] + " " + d[:400]) or CONV_TITLE.search(x["title"]) else ""
    seo = "да" if any(s in p for p in x["paths"] for s in SEO_SECTIONS) or SEO_TITLE.search(x["title"]) else ""
    rows.append([x["id"], x["title"], year.group(1) if year else "", imprint[:150], pg.group(1) if pg else "", lang, textbook, conv, seo,
                 "да" if x["limited"] == "True" else "", ", ".join(x["formats"]), x["size"], x["added"], group(x["page"]), short(x["paths"][0]),
                 "; ".join(short(p) for p in x["paths"][1:]), f"https://www.twirpx.com/file/{x['id']}/"])
rows.sort(key=lambda r: (ROOTS.index(next(k for k, v in group_name.items() if v == r[13])), r[14], r[1].lower()))
head = ["ID", "Название (как на сайте)", "Год", "Выходные данные", "Страниц", "Язык", "Учебник", "Про конверсию", "SEO", "Ограничен (limited)",
        "Форматы", "Размер", "Добавлен", "Группа", "Раздел", "Ещё в разделах", "Ссылка"]
widths = [10, 70, 6, 40, 8, 7, 8, 9, 6, 9, 12, 10, 11, 30, 45, 30, 36]
with open(OUT + ".csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f); w.writerow(head); w.writerows(rows)
wb = openpyxl.Workbook()
def sheet(ws, hdr, data, wd):
    ws.append(hdr)
    for r in data: ws.append(r)
    for i, v in enumerate(wd, 1): ws.column_dimensions[get_column_letter(i)].width = v
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
sheet(wb.active, head, rows, widths); wb.active.title = "Книги"
sheet(wb.create_sheet("Про конверсию"), head, [r for r in rows if r[7]], widths)
sheet(wb.create_sheet("SEO"), head, [r for r in rows if r[8]], widths)
sheet(wb.create_sheet("Удалены правообладателем"), ["ID", "Название (как на сайте)", "Раздел"],
      [[x["id"], x["title"], short(x["paths"][0])] for x in sorted(removed, key=lambda x: (x["paths"][0], x["title"].lower()))], [10, 70, 50])
sheet(wb.create_sheet("Отсеяно — не книги"), ["ID", "Название (как на сайте)", "Почему", "Начало описания", "Раздел", "Ссылка"],
      [[x["id"], x["title"], why, x["description"][:120], short(x["paths"][0]), f"https://www.twirpx.com/file/{x['id']}/"] for x, why in skipped], [10, 60, 18, 60, 40, 36])
tab_names = {"lp": "учебные программы", "standard": "нормативные документы", "article": "статьи и сборники", "paper": "авторефераты и диссертации",
             "guideline": "методички и практикумы", "lecture": "лекции", "package": "УМК", "cheat": "шпаргалки и ответы", "control": "контрольные",
             "course": "курсовые", "degree": "дипломные", "lab": "лабораторные", "presentation": "презентации", "ref": "рефераты и доклады",
             "report": "отчёты по практике", "question": "билеты и вопросы", "dictionary": "словари", "encyclopedia": "энциклопедии", "reference": "справочники"}
base = [p for p in pages if "?ft=" not in p["page"]]
sheet(wb.create_sheet("Разделы"), ["Раздел", "Файлов в списке раздела", "Не взято: вкладки учебных материалов", "Адрес"],
      [[p["name"], p["items"], ", ".join(f"{tab_names.get(t, t)} {n}" for t, n in p.get("tabs", {}).items() if t in NOT_BOOK_TYPES),
        "https://www.twirpx.com" + p["page"]] for p in base], [50, 12, 80, 70])
wb.save(OUT + ".xlsx")
skipped_tabs = sum(n for p in base for t, n in p.get("tabs", {}).items() if t in NOT_BOOK_TYPES)
print("страниц:", len(pages), "| уникальных файлов:", total, "| журналов:", len(journals))
print("удалено правообладателем:", len(removed), "| прочие статусы:", dict(other_status))
print("можно скачать:", len(ok), "| из них не книги:", len(skipped), collections.Counter(w for _, w in skipped))
print("КНИГ В СПИСКЕ:", len(rows), "| рус:", sum(r[5] == "рус" for r in rows), "| учебники:", sum(bool(r[6]) for r in rows),
      "| про конверсию:", sum(bool(r[7]) for r in rows), "| SEO:", sum(bool(r[8]) for r in rows), "| limited:", sum(bool(r[9]) for r in rows))
print("не взято с вкладок учебных материалов:", skipped_tabs)
print("с годом:", sum(bool(r[2]) for r in rows), "| со страницами:", sum(bool(r[4]) for r in rows))
